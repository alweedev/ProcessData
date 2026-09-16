"""Testes do fluxo de substituição de aprovador (Fase 8)."""

import io

import pandas as pd
from _helpers import valid_cpf, xlsx_upload
from openpyxl import load_workbook

OLD = valid_cpf(1)
NEW = valid_cpf(2)
OTHER = valid_cpf(3)


def _users_df(new_status="ATIVO", include_new=True):
    rows = [
        {"CPF": OLD, "Status": "ATIVO", "NomeCompleto": "Aprovador Antigo"},
        {"CPF": OTHER, "Status": "ATIVO", "NomeCompleto": "Outro"},
    ]
    if include_new:
        rows.append({"CPF": NEW, "Status": new_status, "NomeCompleto": "Aprovador Novo"})
    return pd.DataFrame(rows)


def _base_df(rows):
    return pd.DataFrame(rows)


def _post_preview(client, users_df, base_df, cpf, new_cpf, **form):
    data = {
        "users_file": xlsx_upload(users_df, "users.xlsx"),
        "base_file": xlsx_upload(base_df, "base.xlsx"),
        "cpf": cpf,
        "new_cpf": new_cpf,
    }
    data.update({k: str(v) for k, v in form.items()})
    return client.post("/api/aprovacao/substituir/preview", data=data, content_type="multipart/form-data")


def _post_export(client, users_df, base_df, cpf, new_cpf, **form):
    data = {
        "users_file": xlsx_upload(users_df, "users.xlsx"),
        "base_file": xlsx_upload(base_df, "base.xlsx"),
        "cpf": cpf,
        "new_cpf": new_cpf,
        "mode": form.pop("mode", "all"),
    }
    data.update({k: str(v) for k, v in form.items()})
    return client.post("/api/aprovacao/substituir/export", data=data, content_type="multipart/form-data")


def _rows_from_xlsx(resp):
    wb = load_workbook(io.BytesIO(resp.data))
    ws = wb.active
    header = [c.value for c in ws[1]]
    return [dict(zip(header, [c.value for c in row])) for row in ws.iter_rows(min_row=2)]


# ---------------------------------------------------------------- validação


def test_new_cpf_equal_to_old_rejected(client):
    base = _base_df([{"AprovacaoId": "A", "LoginAprovador_1": OLD}])
    resp = _post_preview(client, _users_df(), base, OLD, OLD)
    assert resp.status_code == 400
    assert "diferente" in resp.get_json()["error"].lower()


def test_new_approver_not_found_rejected(client):
    base = _base_df([{"AprovacaoId": "A", "LoginAprovador_1": OLD}])
    resp = _post_preview(client, _users_df(include_new=False), base, OLD, NEW)
    assert resp.status_code == 400
    assert "encontrado" in resp.get_json()["error"].lower()


def test_new_approver_inactive_rejected(client):
    base = _base_df([{"AprovacaoId": "A", "LoginAprovador_1": OLD}])
    resp = _post_preview(client, _users_df(new_status="INATIVO"), base, OLD, NEW)
    assert resp.status_code == 400
    assert "ativo" in resp.get_json()["error"].lower()


def test_old_cpf_not_found_in_any_structure(client):
    base = _base_df([{"AprovacaoId": "A", "LoginAprovador_1": OTHER}])
    resp = _post_export(client, _users_df(), base, OLD, NEW)
    assert resp.status_code == 400
    assert "nenhuma estrutura" in resp.get_json()["error"].lower()


def test_old_approver_not_found_rejected(client):
    base = _base_df([{"AprovacaoId": "A", "LoginAprovador_1": OLD}])
    users = _users_df().drop(index=0).reset_index(drop=True)  # remove OLD da base de usuários
    resp = _post_preview(client, users, base, OLD, NEW)
    assert resp.status_code == 400
    assert "encontrado" in resp.get_json()["error"].lower()


def test_old_approver_inactive_rejected(client):
    base = _base_df([{"AprovacaoId": "A", "LoginAprovador_1": OLD}])
    users = _users_df()
    users.loc[users["CPF"] == OLD, "Status"] = "INATIVO"
    resp = _post_preview(client, users, base, OLD, NEW)
    assert resp.status_code == 400
    assert "ativo" in resp.get_json()["error"].lower()


# ---------------------------------------------------------------- preview


def test_preview_happy_path(client):
    base = _base_df(
        [{"AprovacaoId": "A", "AprovacaoPor": "VIAJANTE", "LoginAprovador_1": OLD, "LoginAprovador_2": OTHER}]
    )
    resp = _post_preview(client, _users_df(), base, OLD, NEW)
    assert resp.status_code == 200, resp.get_data(as_text=True)
    body = resp.get_json()
    assert body["summary"]["estruturasAfetadas"] == 1
    assert body["oldApprover"]["nomeCompleto"] == "Aprovador Antigo"
    assert body["newApprover"]["nomeCompleto"] == "Aprovador Novo"
    assert body["items"][0]["posicoes"] == [1]


def test_preview_flags_duplicate_when_new_already_approver(client):
    base = _base_df([{"AprovacaoId": "A", "LoginAprovador_1": OLD, "LoginAprovador_2": NEW}])
    resp = _post_preview(client, _users_df(), base, OLD, NEW)
    assert resp.status_code == 200, resp.get_data(as_text=True)
    body = resp.get_json()
    assert body["summary"]["estruturasComDuplicidade"] == 1
    assert body["items"][0]["teraDuplicidade"] is True


def test_preview_flags_duplicate_when_new_already_in_second_level(client):
    base = _base_df([{"AprovacaoId": "A", "LoginAprovador_1": OLD, "LoginAprovador_SEGUNDO_NIVEL": NEW}])
    resp = _post_preview(client, _users_df(), base, OLD, NEW)
    assert resp.status_code == 200, resp.get_data(as_text=True)
    body = resp.get_json()
    assert body["summary"]["estruturasComDuplicidade"] == 1
    assert body["items"][0]["teraDuplicidade"] is True


# ---------------------------------------------------------------- export


def test_replace_keeps_same_position(client):
    base = _base_df([{"AprovacaoId": "A", "LoginAprovador_1": OLD, "LoginAprovador_2": OTHER}])
    resp = _post_export(client, _users_df(), base, OLD, NEW)
    assert resp.status_code == 200, resp.get_data(as_text=True)
    rows = _rows_from_xlsx(resp)
    assert len(rows) == 1
    assert str(rows[0]["LoginAprovador_1"] or "").replace(".", "").replace("-", "") == NEW
    assert str(rows[0]["LoginAprovador_2"] or "").replace(".", "").replace("-", "") == OTHER
    assert rows[0]["Operacao"] == "UPDATE"


def test_replace_second_level_only_when_flagged(client):
    base = _base_df([{"AprovacaoId": "A", "LoginAprovador_1": OTHER, "LoginAprovador_SEGUNDO_NIVEL": OLD}])
    resp_off = _post_export(client, _users_df(), base, OLD, NEW, replace_second_level="false")
    assert resp_off.status_code == 200, resp_off.get_data(as_text=True)
    rows_off = _rows_from_xlsx(resp_off)
    assert str(rows_off[0]["LoginAprovador_SEGUNDO_NIVEL"] or "").replace(".", "").replace("-", "") == OLD
    assert (rows_off[0].get("Operacao") or "") == ""

    resp_on = _post_export(client, _users_df(), base, OLD, NEW, replace_second_level="true")
    assert resp_on.status_code == 200, resp_on.get_data(as_text=True)
    rows_on = _rows_from_xlsx(resp_on)
    assert str(rows_on[0]["LoginAprovador_SEGUNDO_NIVEL"] or "").replace(".", "").replace("-", "") == NEW
    assert rows_on[0]["Operacao"] == "UPDATE"


def test_update_only_on_changed_rows(client):
    base = _base_df(
        [
            {"AprovacaoId": "A", "LoginAprovador_1": OLD},
            {"AprovacaoId": "A", "LoginAprovador_1": OTHER},
        ]
    )
    resp = _post_export(client, _users_df(), base, OLD, NEW)
    assert resp.status_code == 200, resp.get_data(as_text=True)
    rows = _rows_from_xlsx(resp)
    assert len(rows) == 2
    ops = sorted((r["Operacao"] or "") for r in rows)
    assert ops == ["", "UPDATE"]


def test_duplicate_gate_blocks_without_flag(client):
    base = _base_df([{"AprovacaoId": "A", "LoginAprovador_1": OLD, "LoginAprovador_2": NEW}])
    resp = _post_export(client, _users_df(), base, OLD, NEW)
    assert resp.status_code == 400
    body = resp.get_json()
    assert body.get("warning") is True
    assert body.get("estruturasComDuplicidade")


def test_duplicate_gate_bypassed_with_flag(client):
    base = _base_df([{"AprovacaoId": "A", "LoginAprovador_1": OLD, "LoginAprovador_2": NEW}])
    resp = _post_export(client, _users_df(), base, OLD, NEW, ignore_duplicate_warning="true")
    assert resp.status_code == 200, resp.get_data(as_text=True)
    rows = _rows_from_xlsx(resp)
    assert str(rows[0]["LoginAprovador_1"] or "").replace(".", "").replace("-", "") == NEW
    assert str(rows[0]["LoginAprovador_2"] or "").replace(".", "").replace("-", "") == NEW


def test_mode_selected_restricts_target_structures(client):
    base = _base_df(
        [
            {"AprovacaoId": "A", "LoginAprovador_1": OLD},
            {"AprovacaoId": "B", "LoginAprovador_1": OLD},
        ]
    )
    resp = _post_export(client, _users_df(), base, OLD, NEW, mode="selected", **{"selected_aprovacao_ids": "A"})
    assert resp.status_code == 200, resp.get_data(as_text=True)
    rows = _rows_from_xlsx(resp)
    assert len(rows) == 1
    assert rows[0]["AprovacaoId"] == "A"
