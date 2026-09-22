import io
import json
import os
import zipfile

import pandas as pd
from _cascade_fixtures import USR_A, A, B, C, cad, est, viajante
from _helpers import xlsx_upload
from openpyxl import load_workbook

from backend.core.config import settings
from backend.services.audit_service import AuditService
from backend.services.inactivation_cascade_service import InactivationCascadeService
from backend.shared.cpf_mask import mascarar_cpf

ROTA = "/api/inativacao/executar"


def _digital(cadastro, estruturas, cpfs):
    return InactivationCascadeService.analisar(cadastro, estruturas, cpfs).payload["impressaoDigital"]


def _post(client, cadastro, estruturas, cpfs, digital=None, **campos):
    data = {
        "cadastro": xlsx_upload(cadastro, "cadastro.xlsx"),
        "estruturas": xlsx_upload(estruturas, "estruturas.xlsx"),
        "cpfs": json.dumps(cpfs),
        "impressaoDigital": digital if digital is not None else _digital(cadastro, estruturas, cpfs),
    }
    data.update(campos)
    return client.post(ROTA, data=data, content_type="multipart/form-data")


def _planilha(zf, nome):
    return pd.read_excel(io.BytesIO(zf.read(nome)), dtype=str).fillna("")


def test_devolve_zip_com_as_duas_planilhas(client):
    cadastro, estruturas = cad(USR_A), est(viajante("S1", A, B), viajante("S2", C, B, A))
    resp = _post(client, cadastro, estruturas, [A])
    assert resp.status_code == 200, resp.get_data(as_text=True)
    assert resp.headers["Content-Type"].startswith("application/zip")
    with zipfile.ZipFile(io.BytesIO(resp.data)) as zf:
        assert zf.namelist() == ["saida_inativacao.xlsx", "estruturas_atualizadas.xlsx"]
        assert _planilha(zf, "saida_inativacao.xlsx")["Operacao"].tolist() == ["DELETE"]
        est_saida = _planilha(zf, "estruturas_atualizadas.xlsx")
        assert dict(zip(est_saida["AprovacaoId"], est_saida["Operacao"])) == {"S1": "DELETE", "S2": "UPDATE"}


def test_grava_auditoria_sem_dado_pessoal(client):
    resp = _post(client, cad(USR_A), est(viajante("S1", A, B)), [A])
    assert resp.status_code == 200
    eventos = AuditService.list_events()
    assert any(e["event_type"] == "inativacao_execucao" and e["status"] == "success" for e in eventos)
    texto = json.dumps(eventos, ensure_ascii=False)
    for pessoal in (A, "Ana Souza", "ana@x.com"):
        assert pessoal not in texto


def test_impressao_digital_divergente_da_409(client):
    resp = _post(client, cad(USR_A), est(viajante("S1", A, B)), [A], digital="0" * 64)
    assert resp.status_code == 409
    assert resp.get_json()["code"] == "ANALISE_DIVERGENTE"


def test_estrutura_orfa_pede_confirmacao_e_depois_executa(client):
    cadastro, estruturas = cad(USR_A), est(viajante("S3", C, A))
    resp = _post(client, cadastro, estruturas, [A])
    assert resp.status_code == 400
    corpo = resp.get_json()
    assert corpo["code"] == "ORFAS_SEM_CONFIRMACAO"
    assert corpo["estruturasOrfas"] == ["S3"]
    assert _post(client, cadastro, estruturas, [A], ignore_orphan_warning="true").status_code == 200


def test_nada_a_executar(client):
    resp = _post(client, cad(USR_A), est(viajante("S1", C, B)), [], digital="x")
    assert resp.status_code == 400
    assert resp.get_json()["code"] == "NADA_A_EXECUTAR"


def test_sem_arquivos_nunca_responde_sucesso(client):
    resp = client.post(ROTA, data={"cpfs": json.dumps([A])}, content_type="multipart/form-data")
    assert resp.status_code == 400
    assert resp.get_json()["code"] == "BASE_AUSENTE"


def test_arquivo_acima_do_teto_da_rota(client, monkeypatch):
    monkeypatch.setattr(settings, "INATIVACAO_MAX_CONTENT_LENGTH", 200, raising=False)
    resp = _post(client, cad(USR_A), est(viajante("S1", A, B)), [A])
    assert resp.status_code == 413
    assert resp.get_json()["code"] == "ARQUIVO_GRANDE"


def test_apaga_os_temporarios(client):
    _post(client, cad(USR_A), est(viajante("S1", A, B)), [A])
    _post(client, cad(USR_A), est(viajante("S1", A, B)), [A], digital="0" * 64)
    assert os.listdir(settings.UPLOAD_FOLDER) == []


def test_erro_interno_vira_500_e_nunca_sucesso(client, monkeypatch):
    def boom(*args, **kwargs):
        raise RuntimeError("bug interno")

    monkeypatch.setattr(InactivationCascadeService, "executar", boom)
    resp = _post(client, cad(USR_A), est(viajante("S1", A, B)), [A], digital="x")
    assert resp.status_code == 500
    assert resp.get_json()["code"] == "ERRO_INTERNO"


def test_rotas_antigas_foram_removidas(client):
    for rota in ("/api/process_inativacao", "/api/preview_inativacao", "/api/inativacao/buscar"):
        # Rota removida: 404 (sem regra) ou 405 (só o curinga do SPA aceita GET); nunca 200/400/500.
        assert client.post(rota, data={}, content_type="multipart/form-data").status_code in (404, 405)


def test_usuario_sem_estruturas_gera_planilha_de_estruturas_valida_so_com_cabecalho(client):
    resp = _post(client, cad(USR_A), est(viajante("S1", C, B)), [A])
    assert resp.status_code == 200, resp.get_data(as_text=True)
    with zipfile.ZipFile(io.BytesIO(resp.data)) as zf:
        wb = load_workbook(io.BytesIO(zf.read("estruturas_atualizadas.xlsx")))
    linhas = list(wb["Aprovacao"].iter_rows(values_only=True))
    assert len(linhas) == 1
    assert linhas[0][:2] == ("Operacao", "AprovacaoId")


def test_estrutura_compartilhada_vira_400_com_codigo_proprio(client):
    estruturas = est(viajante("S1", A, B), viajante("S1", C, B))
    data = {
        "cadastro": xlsx_upload(cad(USR_A), "cadastro.xlsx"),
        "estruturas": xlsx_upload(estruturas, "estruturas.xlsx"),
        "cpfs": json.dumps([A]),
        "impressaoDigital": "x",
    }
    resp = client.post(ROTA, data=data, content_type="multipart/form-data")
    assert resp.status_code == 400
    assert resp.get_json()["code"] == "ESTRUTURA_COMPARTILHADA"


def test_auditoria_lista_so_os_cpfs_executados_e_nao_os_enviados_pelo_cliente(client, monkeypatch):
    cadastro, estruturas = cad(USR_A), est(viajante("S1", A, B))
    real = InactivationCascadeService.executar

    def executa_so_o_cpf_valido(df_cadastro, df_estruturas, cpfs, digital, ignorar=False, confirmar=False):
        return real(df_cadastro, df_estruturas, [A], _digital(df_cadastro, df_estruturas, [A]), ignorar, confirmar)

    monkeypatch.setattr(InactivationCascadeService, "executar", staticmethod(executa_so_o_cpf_valido))
    assert _post(client, cadastro, estruturas, [A, C], digital="x").status_code == 200
    evento = next(e for e in AuditService.list_events() if e["event_type"] == "inativacao_execucao")
    assert evento["details"]["cpfs"] == [mascarar_cpf(A)]
