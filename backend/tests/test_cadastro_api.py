"""Testes do endpoint de cadastro (/api/process_cadastro).

Inclui as verificações portadas do antigo script backend/test_cadastro_run.py.
"""

import io

import pandas as pd
from _helpers import valid_cpf, xlsx_upload
from openpyxl import load_workbook

from backend.services.processing_service import ProcessingService


def _read_first_sheet(binary):
    ws = load_workbook(io.BytesIO(binary)).active
    header = [c.value for c in ws[1]]
    return header, [dict(zip(header, [c.value for c in row])) for row in ws.iter_rows(min_row=2)]


def test_docx_upload_rejected(client):
    data = {"files[]": (io.BytesIO(b"PK\x03\x04 fake docx"), "ficha.docx")}
    resp = client.post("/api/process_cadastro", data=data, content_type="multipart/form-data")
    assert resp.status_code == 400
    assert "extens" in resp.get_json()["error"].lower()


def test_cadastro_happy_path_returns_xlsx(client):
    df = pd.DataFrame(
        [
            {
                "CPF": valid_cpf(10),
                "NOME COMPLETO": "Ana Souza",
                "EMAIL": "ana@empresa.com",
                "EMPRESA": "Empresa A",
                "Centro de custo": "CC1",
                "SOLICITANTE? (S/N)": "S",
            },
        ]
    )
    data = {"files[]": xlsx_upload(df, "cadastro.xlsx"), "login_choice": "CPF", "fluxo": "SELF"}
    resp = client.post("/api/process_cadastro", data=data, content_type="multipart/form-data")
    assert resp.status_code == 200, resp.get_data(as_text=True)
    assert "spreadsheetml" in resp.headers.get("Content-Type", "")
    assert resp.data[:2] == b"PK"  # xlsx é um zip

    header, rows = _read_first_sheet(resp.data)
    assert rows[0]["Solicitante"] == "S"
    # Login = CPF formatado (XXXXXXXXX-XX) quando login_choice=CPF
    assert str(rows[0]["Login"]).replace(".", "").replace("-", "") == valid_cpf(10)
    assert "-" in str(rows[0]["Login"])


def test_cadastro_success_audit_tracks_invalid_row_count_without_raw_messages(client):
    """O histórico registra QUANTAS linhas falharam (métrica), mas não o
    conteúdo bruto de `errors` -- /api/history é legível sem autenticação
    (GET), então mensagens de erro (que em outros fluxos podem conter
    caminho de arquivo do servidor) não devem ficar ali."""
    from backend.services.audit_service import AuditService

    df = pd.DataFrame(
        [
            {
                "CPF": valid_cpf(11),
                "NOME COMPLETO": "Carlos Teste",
                "EMAIL": "carlos-sem-dominio-valido",
                "EMPRESA": "Empresa A",
                "Centro de custo": "CC1",
                "SOLICITANTE? (S/N)": "S",
            },
        ]
    )
    data = {"files[]": xlsx_upload(df, "cadastro.xlsx"), "login_choice": "CPF", "fluxo": "SELF"}
    resp = client.post("/api/process_cadastro", data=data, content_type="multipart/form-data")
    assert resp.status_code == 200, resp.get_data(as_text=True)

    events = AuditService.list_events(limit=5)
    latest = next(e for e in events if e["event_type"] == "cadastro" and e["status"] == "success")
    assert latest["details"]["invalid_rows"] >= 1
    assert "errors" not in latest["details"]


def test_cadastro_multi_file_upload_merges_rows(client):
    """/api/process_cadastro nunca era testado com mais de 1 arquivo, embora
    a rota faça `request.files.getlist('files[]')` e itere sobre todos."""
    df1 = pd.DataFrame(
        [
            {
                "CPF": valid_cpf(20),
                "NOME COMPLETO": "Ana Primeiro Arquivo",
                "EMAIL": "ana@empresa.com",
                "EMPRESA": "Empresa A",
                "Centro de custo": "CC1",
                "SOLICITANTE? (S/N)": "S",
            },
        ]
    )
    df2 = pd.DataFrame(
        [
            {
                "CPF": valid_cpf(21),
                "NOME COMPLETO": "Bruno Segundo Arquivo",
                "EMAIL": "bruno@empresa.com",
                "EMPRESA": "Empresa B",
                "Centro de custo": "CC2",
                "SOLICITANTE? (S/N)": "N",
            },
        ]
    )
    data = {
        "files[]": [xlsx_upload(df1, "um.xlsx"), xlsx_upload(df2, "dois.xlsx")],
        "login_choice": "CPF",
        "fluxo": "SELF",
    }
    resp = client.post("/api/process_cadastro", data=data, content_type="multipart/form-data")
    assert resp.status_code == 200, resp.get_data(as_text=True)

    _header, rows = _read_first_sheet(resp.data)
    assert len(rows) == 2
    nomes = {r["NomeCompleto"] for r in rows}
    assert nomes == {"ANA PRIMEIRO ARQUIVO", "BRUNO SEGUNDO ARQUIVO"}


def test_cadastro_front_flow_via_endpoint(client):
    """fluxo=FRONT só era testado direto no ProcessingService; aqui confere
    que o endpoint HTTP repassa `fluxo` corretamente."""
    df = pd.DataFrame(
        [
            {
                "EMAIL": "viajante@empresa.com",
                "NOME COMPLETO": "Viajante Teste",
                "EMPRESA": "Empresa A",
                "Centro de custo": "CC1",
                "SOLICITANTE? (S/N)": "N",
            },
        ]
    )
    data = {"files[]": xlsx_upload(df, "front.xlsx"), "login_choice": "EMAIL", "fluxo": "FRONT"}
    resp = client.post("/api/process_cadastro", data=data, content_type="multipart/form-data")
    assert resp.status_code == 200, resp.get_data(as_text=True)

    _header, rows = _read_first_sheet(resp.data)
    assert rows[0]["Login"] == "FRONTVIAJANTE@EMPRESA.COM"
    assert rows[0]["ViajanteMasterNacional"] == "S"
    assert rows[0]["ViajanteMasterInternacional"] == "S"


def test_cadastro_bool_and_blank_rows(tmp_path):
    df = pd.DataFrame(
        [
            {
                "CPF": valid_cpf(1),
                "NOME COMPLETO": "Ana Um",
                "EMAIL": "a@x.com",
                "EMPRESA": "E",
                "SOLICITANTE? (S/N)": "Sim",
                "Terceiro": "Sim",
            },
            {
                "CPF": valid_cpf(2),
                "NOME COMPLETO": "Bruno Dois",
                "EMAIL": "b@x.com",
                "EMPRESA": "E",
                "SOLICITANTE? (S/N)": "",
                "Terceiro": "Não",
            },
            {"CPF": "", "NOME COMPLETO": "", "EMAIL": "", "EMPRESA": "", "SOLICITANTE? (S/N)": ""},
        ]
    )
    p = tmp_path / "c.xlsx"
    df.to_excel(p, index=False)
    _errors, out = ProcessingService.process_records_from_files([str(p)], login_choice="CPF", fluxo="SELF")

    assert len(out) == 2  # linha totalmente em branco descartada
    assert set(out["Solicitante"].unique()).issubset({"S", "N"})
    assert out["Solicitante"].tolist() == ["S", "N"]
    assert out["Terceiro"].tolist() == ["S", "N"]
