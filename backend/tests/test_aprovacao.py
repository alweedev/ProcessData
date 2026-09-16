"""Testes de validação e remoção de aprovador (B3: P4 + P5)."""

import io

import pandas as pd
from _helpers import valid_cpf, xlsx_bytes, xlsx_upload
from openpyxl import load_workbook

APPROVER = valid_cpf(1)
SECOND = valid_cpf(2)
OTHER = valid_cpf(3)


def _users_df(status="ATIVO", with_name=True, include_approver=True):
    cols = {"CPF": [], "Status": []}
    if with_name:
        cols["NomeCompleto"] = []
    rows = []
    if include_approver:
        row = {"CPF": APPROVER, "Status": status}
        if with_name:
            row["NomeCompleto"] = "Aprovador Um"
        rows.append(row)
    rows.append({"CPF": OTHER, "Status": "ATIVO", **({"NomeCompleto": "Outro"} if with_name else {})})
    return pd.DataFrame(rows)


def _base_df(rows):
    return pd.DataFrame(rows)


def _post_export(client, users_df, base_df, cpf, **form):
    data = {
        "users_file": xlsx_upload(users_df, "users.xlsx"),
        "base_file": xlsx_upload(base_df, "base.xlsx"),
        "cpf": cpf,
        "mode": form.pop("mode", "all"),
    }
    data.update({k: str(v) for k, v in form.items()})
    return client.post("/api/aprovacao/remover/export", data=data, content_type="multipart/form-data")


def _post_preview(client, users_df, base_df, cpf, **form):
    data = {
        "users_file": xlsx_upload(users_df, "users.xlsx"),
        "base_file": xlsx_upload(base_df, "base.xlsx"),
        "cpf": cpf,
    }
    data.update({k: str(v) for k, v in form.items()})
    return client.post("/api/aprovacao/remover/preview", data=data, content_type="multipart/form-data")


# ---------------------------------------------------------------- P4


def test_cpf_checksum_rejected(client):
    bad = "12345678900"  # 11 dígitos, dígito verificador inválido
    resp = _post_preview(client, _users_df(), _base_df([{"AprovacaoId": "A", "LoginAprovador_1": APPROVER}]), bad)
    assert resp.status_code == 400
    assert "verificador" in resp.get_json()["error"].lower()


def test_empty_cpf_rejected(client):
    resp = _post_preview(client, _users_df(), _base_df([{"AprovacaoId": "A", "LoginAprovador_1": APPROVER}]), "")
    assert resp.status_code == 400
    assert "informe um cpf" in resp.get_json()["error"].lower()


def test_cpf_too_long_rejected(client):
    # normalize_cpf_input faz zfill(11) antes de checar o tamanho: uma
    # entrada curta (poucos dígitos) sempre vira 11 dígitos e cai no erro de
    # dígito verificador (test_cpf_checksum_rejected), não no de tamanho.
    # Só dá pra disparar "CPF inválido. Informe 11 dígitos." com excesso.
    resp = _post_preview(
        client, _users_df(), _base_df([{"AprovacaoId": "A", "LoginAprovador_1": APPROVER}]), "123456789012"
    )
    assert resp.status_code == 400
    assert "11 dígitos" in resp.get_json()["error"]


def test_users_base_missing_cpf_column(client):
    users = pd.DataFrame([{"NomeCompleto": "Aprovador Um", "Status": "ATIVO"}])
    base = _base_df([{"AprovacaoId": "A", "LoginAprovador_1": APPROVER}])
    resp = _post_preview(client, users, base, APPROVER)
    assert resp.status_code == 400
    assert "cpf" in resp.get_json()["error"].lower()


def test_users_base_missing_name_column(client):
    users = _users_df(with_name=False)
    base = _base_df([{"AprovacaoId": "A", "LoginAprovador_1": APPROVER}])
    resp = _post_preview(client, users, base, APPROVER)
    assert resp.status_code == 400
    assert "nome" in resp.get_json()["error"].lower()


def test_approver_inactive_rejected(client):
    users = _users_df(status="INATIVO")
    base = _base_df([{"AprovacaoId": "A", "LoginAprovador_1": APPROVER}])
    resp = _post_preview(client, users, base, APPROVER)
    assert resp.status_code == 400
    assert "ativo" in resp.get_json()["error"].lower()


def test_approver_not_found(client):
    users = _users_df(include_approver=False)
    base = _base_df([{"AprovacaoId": "A", "LoginAprovador_1": APPROVER}])
    resp = _post_preview(client, users, base, APPROVER)
    assert resp.status_code == 400
    assert "encontrado" in resp.get_json()["error"].lower()


def test_preview_happy_path(client):
    base = _base_df(
        [
            {"AprovacaoId": "A", "AprovacaoPor": "VIAJANTE", "LoginAprovador_1": APPROVER, "LoginAprovador_2": OTHER},
        ]
    )
    resp = _post_preview(client, _users_df(), base, APPROVER)
    assert resp.status_code == 200, resp.get_data(as_text=True)
    body = resp.get_json()
    assert body["summary"]["estruturasAfetadas"] >= 1
    assert body["approver"]["nomeCompleto"] == "Aprovador Um"


# ---------------------------------------------------------------- P5


def test_second_level_promoted(client):
    base = _base_df(
        [
            {
                "AprovacaoId": "A",
                "LoginAprovador_1": APPROVER,
                "LoginAprovador_2": "",
                "LoginAprovador_SEGUNDO_NIVEL": SECOND,
            }
        ]
    )
    resp = _post_export(client, _users_df(), base, APPROVER, remove_second_level="false")
    assert resp.status_code == 200, resp.get_data(as_text=True)
    wb = load_workbook(io.BytesIO(resp.data))
    ws = wb.active
    header = [c.value for c in ws[1]]
    data = [dict(zip(header, [c.value for c in row])) for row in ws.iter_rows(min_row=2)]
    assert len(data) == 1
    assert str(data[0]["LoginAprovador_1"] or "").replace(".", "").replace("-", "") == SECOND
    assert (data[0].get("LoginAprovador_SEGUNDO_NIVEL") or "") == ""
    assert data[0]["Operacao"] == "UPDATE"


def test_second_level_not_promoted_when_same_cpf_being_removed(client):
    """Se o próprio CPF removido também está no segundo nível (mantido por
    remove_second_level=false), ele não pode ser promovido de volta ao 1º
    nível — isso reintroduziria o aprovador que a operação removeu."""
    base = _base_df(
        [
            {
                "AprovacaoId": "A",
                "LoginAprovador_1": APPROVER,
                "LoginAprovador_2": "",
                "LoginAprovador_SEGUNDO_NIVEL": APPROVER,
            }
        ]
    )
    resp = _post_export(client, _users_df(), base, APPROVER, remove_second_level="false")
    assert resp.status_code == 200, resp.get_data(as_text=True)
    wb = load_workbook(io.BytesIO(resp.data))
    ws = wb.active
    header = [c.value for c in ws[1]]
    data = [dict(zip(header, [c.value for c in row])) for row in ws.iter_rows(min_row=2)]
    assert len(data) == 1
    assert str(data[0]["LoginAprovador_1"] or "") == ""
    assert str(data[0]["LoginAprovador_SEGUNDO_NIVEL"] or "").replace(".", "").replace("-", "") == APPROVER


def test_update_only_on_changed_rows(client):
    base = _base_df(
        [
            {"AprovacaoId": "A", "LoginAprovador_1": APPROVER, "LoginAprovador_2": OTHER},
            {"AprovacaoId": "A", "LoginAprovador_1": OTHER, "LoginAprovador_2": ""},
        ]
    )
    resp = _post_export(client, _users_df(), base, APPROVER, ignore_empty_warning="true")
    assert resp.status_code == 200, resp.get_data(as_text=True)
    wb = load_workbook(io.BytesIO(resp.data))
    ws = wb.active
    header = [c.value for c in ws[1]]
    rows = [dict(zip(header, [c.value for c in row])) for row in ws.iter_rows(min_row=2)]
    assert len(rows) == 2
    ops = sorted((r["Operacao"] or "") for r in rows)
    assert ops == ["", "UPDATE"]


def test_structure_truly_empty_triggers_gate(client):
    base = _base_df([{"AprovacaoId": "A", "LoginAprovador_1": APPROVER, "LoginAprovador_2": ""}])
    resp = _post_export(client, _users_df(), base, APPROVER)  # sem ignore_empty_warning
    assert resp.status_code == 400
    body = resp.get_json()
    assert body.get("warning") is True
    assert body.get("estruturasSemAprovador")


def test_second_level_removal_can_also_trigger_empty_gate(client):
    """O gate de estrutura vazia deve considerar a remoção do 2º nível
    também, não só do 1º (remove_second_level=true esvaziando o único
    fallback que restava)."""
    base = _base_df([{"AprovacaoId": "A", "LoginAprovador_1": APPROVER, "LoginAprovador_SEGUNDO_NIVEL": APPROVER}])
    resp = _post_export(client, _users_df(), base, APPROVER, remove_second_level="true")
    assert resp.status_code == 400
    body = resp.get_json()
    assert body.get("warning") is True
    assert body.get("estruturasSemAprovador")


def test_mode_selected_restricts_target_structures(client):
    base = _base_df(
        [
            {"AprovacaoId": "A", "LoginAprovador_1": APPROVER, "LoginAprovador_2": OTHER},
            {"AprovacaoId": "B", "LoginAprovador_1": APPROVER, "LoginAprovador_2": OTHER},
        ]
    )
    resp = _post_export(client, _users_df(), base, APPROVER, mode="selected", **{"selected_aprovacao_ids": "A"})
    assert resp.status_code == 200, resp.get_data(as_text=True)
    wb = load_workbook(io.BytesIO(resp.data))
    ws = wb.active
    header = [c.value for c in ws[1]]
    rows = [dict(zip(header, [c.value for c in row])) for row in ws.iter_rows(min_row=2)]
    assert len(rows) == 1
    assert rows[0]["AprovacaoId"] == "A"


def test_mode_selected_without_ids_rejected(client):
    base = _base_df([{"AprovacaoId": "A", "LoginAprovador_1": APPROVER, "LoginAprovador_2": OTHER}])
    resp = _post_export(client, _users_df(), base, APPROVER, mode="selected")
    assert resp.status_code == 400
    assert "selected_aprovacao_ids" in resp.get_json()["error"]
