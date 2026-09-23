"""Neutralização de formula injection no Excel exportado (P8a)."""

import io

import pandas as pd
from _helpers import valid_cpf, xlsx_upload
from openpyxl import load_workbook

from backend.services.export_service import ExportService

PAYLOADS = ["=SUM(A1)", "+1", "-1", "@x", "\tx"]
_STRING_TYPES = {"s", "str", "inlineStr"}  # qualquer um != fórmula


def test_formula_cells_written_as_text():
    out = ExportService.to_excel_bytes(pd.DataFrame({"x": PAYLOADS}), sheet_name="S")
    ws = load_workbook(io.BytesIO(out.getvalue()))["S"]
    for i, expected in enumerate(PAYLOADS, start=2):  # linha 1 = cabeçalho
        cell = ws.cell(row=i, column=1)
        assert cell.data_type in _STRING_TYPES, f"{expected!r} veio como {cell.data_type}"
        assert cell.data_type != "f"
        assert cell.value == expected


def test_normal_values_unchanged():
    out = ExportService.to_excel_bytes(pd.DataFrame({"x": ["ABC", "Joao", "123"]}), sheet_name="S")
    ws = load_workbook(io.BytesIO(out.getvalue()))["S"]
    assert ws.cell(row=2, column=1).value == "ABC"


def test_happy_path_applies_header_style_and_layout():
    out = ExportService.to_excel_bytes(pd.DataFrame({"NomeCompleto": ["Ana", "Bruno"]}), sheet_name="S")
    ws = load_workbook(io.BytesIO(out.getvalue()))["S"]
    header_cell = ws.cell(row=1, column=1)
    assert header_cell.font.bold is True
    assert header_cell.fill.fgColor.rgb == "00FFDCE6" or header_cell.fill.fgColor.rgb == "FFDCE6F1"
    assert ws.freeze_panes == "A2"
    assert ws.column_dimensions["A"].width is not None
    assert ws.auto_filter.ref is not None


def test_fallback_without_style_still_neutralizes_formulas(monkeypatch):
    """Se a etapa de estilo falhar (ex.: incompatibilidade de versão do
    openpyxl), o export cai no fallback sem estilo -- mas a neutralização de
    formula-injection (a parte que importa pra segurança) precisa continuar
    funcionando mesmo nesse caminho, hoje sem nenhum teste cobrindo isso."""

    def boom(*_a, **_k):
        raise RuntimeError("estilo indisponível")

    # get_column_letter só é usado pela etapa de largura de coluna do
    # ExportService (depois que df.to_excel/_neutralize_worksheet já
    # rodaram) -- ao contrário de Font/PatternFill, não é usado pelo
    # pipeline interno do pandas, então não derruba o próprio fallback.
    monkeypatch.setattr("openpyxl.utils.get_column_letter", boom)

    out = ExportService.to_excel_bytes(pd.DataFrame({"x": ["=SUM(A1)", "normal"]}), sheet_name="S")
    ws = load_workbook(io.BytesIO(out.getvalue()))["S"]
    cell = ws.cell(row=2, column=1)
    assert cell.data_type in _STRING_TYPES
    assert cell.value == "=SUM(A1)"
    # sem estilo: freeze_panes só é setado depois da etapa que forçamos a
    # falhar, então no fallback ele fica ausente (diferente do caminho feliz,
    # coberto por test_happy_path_applies_header_style_and_layout).
    assert ws.freeze_panes is None


def test_aprovacao_export_neutralizes_injection(client):
    users = pd.DataFrame([{"CPF": valid_cpf(1), "NomeCompleto": "Aprovador Um", "Status": "ATIVO"}])
    # payload com prefixo "+": sobrevive ao round-trip de leitura da planilha
    # como string (um "=" seria lido como fórmula sem valor em cache).
    base = pd.DataFrame(
        [
            {
                "AprovacaoId": "A",
                "AprovacaoPor": "+cmd|' /c calc'!A0",
                "LoginAprovador_1": valid_cpf(1),
                "LoginAprovador_2": valid_cpf(3),
            }
        ]
    )
    data = {
        "users_file": xlsx_upload(users, "u.xlsx"),
        "base_file": xlsx_upload(base, "b.xlsx"),
        "cpf": valid_cpf(1),
        "mode": "all",
        "ignore_empty_warning": "true",
    }
    resp = client.post("/api/aprovacao/remover/export", data=data, content_type="multipart/form-data")
    assert resp.status_code == 200, resp.get_data(as_text=True)
    ws = load_workbook(io.BytesIO(resp.data)).active
    header = [c.value for c in ws[1]]
    col = header.index("AprovacaoPor") + 1
    cell = ws.cell(row=2, column=col)
    assert cell.data_type in _STRING_TYPES  # nunca "f" (fórmula)
    assert str(cell.value).startswith("+cmd")


def test_dataframe_vazio_gera_xlsx_valido_com_o_cabecalho():
    out = ExportService.to_excel_bytes(pd.DataFrame(columns=["Operacao", "AprovacaoId"]), sheet_name="S")
    ws = load_workbook(io.BytesIO(out.getvalue()))["S"]
    assert [c.value for c in ws[1]] == ["Operacao", "AprovacaoId"]
    assert ws.max_row == 1
