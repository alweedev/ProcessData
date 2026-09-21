import io
import zipfile

import pandas as pd

# Caracteres que, no início de uma célula, o Excel/Sheets pode interpretar como
# fórmula ("formula injection" / CSV injection). Células assim são forçadas a
# texto explícito no arquivo — nada visível muda para quem abre a planilha.
_RISKY_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def _neutralize_worksheet(ws) -> None:
    for row in ws.iter_rows():
        for cell in row:
            value = cell.value
            if isinstance(value, str) and value.startswith(_RISKY_PREFIXES):
                # data_type "s": openpyxl grava como string, não como <f>órmula
                cell.data_type = "s"


class ExportService:
    @staticmethod
    def to_excel_bytes(df: pd.DataFrame, sheet_name: str = "Dados") -> io.BytesIO:
        output = io.BytesIO()
        try:
            from openpyxl.styles import Alignment, Font, PatternFill
            from openpyxl.utils import get_column_letter

            with pd.ExcelWriter(output, engine="openpyxl") as writer:
                df.to_excel(writer, sheet_name=sheet_name, index=False)
                ws = writer.sheets[sheet_name]

                _neutralize_worksheet(ws)

                header_fill = PatternFill(
                    start_color="FFDCE6F1",
                    end_color="FFDCE6F1",
                    fill_type="solid",
                )
                for cell in list(ws[1]):
                    cell.font = Font(bold=True)
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.fill = header_fill

                for idx, col in enumerate(df.columns, 1):
                    series = df[col].astype(str).fillna("")
                    max_len = max(series.map(len).max(), len(str(col))) + 2
                    max_len = min(max_len, 60)
                    ws.column_dimensions[get_column_letter(idx)].width = max_len

                ws.freeze_panes = "A2"
                try:
                    ws.auto_filter.ref = ws.dimensions
                except Exception:
                    pass

            output.seek(0)
        except Exception:
            # Fallback sem estilo, mas ainda neutralizando fórmulas.
            output = io.BytesIO()
            with pd.ExcelWriter(output, engine="openpyxl") as writer:
                df.to_excel(writer, sheet_name=sheet_name, index=False)
                _neutralize_worksheet(writer.sheets[sheet_name])
            output.seek(0)
        return output

    @staticmethod
    def to_zip_bytes(files: dict[str, io.BytesIO]) -> io.BytesIO:
        """Empacota arquivos já gerados (nome -> bytes) num ZIP em memória."""
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as zf:
            for name, data in files.items():
                zf.writestr(name, data.getvalue())
        output.seek(0)
        return output
